# ============================================================
# dataset/maintain_dataset.py
#
# Dataset maintenance utility — fixes inconsistencies between
# input_questions.csv and the output Excel file.
#
# Usage:
#   # Show current status:
#   python dataset/maintain_dataset.py --status
#
#   # Reset failed rows in both CSV and Excel:
#   python dataset/maintain_dataset.py --fix-failed
#
#   # Reset specific Q_IDs:
#   python dataset/maintain_dataset.py --reset E-DE-010 E-DE-011
#
#   # Reset entire category:
#   python dataset/maintain_dataset.py --reset-category easy_deep_learning
#
#   # Sync CSV with Excel (marks done in CSV if in Excel):
#   python dataset/maintain_dataset.py --sync
#
# Author: Mohammed Danish Alam
# Project: P000233DS — EdGenAI Technologies Capstone
# ============================================================

import os
import csv
import argparse
from collections import Counter

import openpyxl

# ── Paths ─────────────────────────────────────────────────
_DIR        = os.path.dirname(os.path.abspath(__file__))
INPUT_CSV   = os.path.join(_DIR, "inputs", "input_questions.csv")
OUTPUT_XLSX = os.path.join(_DIR, "outputs",
                           "P000233DS_Question_Dataset.xlsx")


# ════════════════════════════════════════════════════════════
# CSV helpers
# ════════════════════════════════════════════════════════════

def load_csv() -> tuple:
    with open(INPUT_CSV, "r", encoding="utf-8") as f:
        reader     = csv.DictReader(f)
        fieldnames = reader.fieldnames
        rows       = list(reader)
    return rows, fieldnames


def save_csv(rows: list, fieldnames: list):
    with open(INPUT_CSV, "w", newline="",
              encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


# ════════════════════════════════════════════════════════════
# Excel helpers
# ════════════════════════════════════════════════════════════

def load_excel_ids() -> dict:
    """Return {q_id: status} from Excel."""
    if not os.path.exists(OUTPUT_XLSX):
        return {}
    wb  = openpyxl.load_workbook(OUTPUT_XLSX)
    ws  = wb.active
    ids = {}
    for row in ws.iter_rows(min_row=3, values_only=True):
        if row[0]:
            ids[str(row[0])] = str(row[19]) \
                if row[19] else ""
    return ids


def delete_from_excel(q_ids: list):
    """Delete rows from Excel by Q_ID."""
    if not os.path.exists(OUTPUT_XLSX):
        print("[Excel] File not found — skipping")
        return 0

    wb      = openpyxl.load_workbook(OUTPUT_XLSX)
    ws      = wb.active
    to_del  = []

    for row in ws.iter_rows(min_row=3):
        if row[0].value and str(row[0].value) in q_ids:
            to_del.append(row[0].row)

    for row_num in sorted(to_del, reverse=True):
        ws.delete_rows(row_num)

    wb.save(OUTPUT_XLSX)
    return len(to_del)


# ════════════════════════════════════════════════════════════
# Commands
# ════════════════════════════════════════════════════════════

def cmd_status():
    """Show current dataset status."""
    rows, _    = load_csv()
    excel_ids  = load_excel_ids()

    csv_counts = Counter(r["status"] for r in rows)
    print("\n" + "=" * 55)
    print("DATASET STATUS")
    print("=" * 55)
    print(f"\n  CSV ({INPUT_CSV}):")
    print(f"    Total:   {len(rows)}")
    print(f"    Pending: {csv_counts.get('pending', 0)}")
    print(f"    Done:    {csv_counts.get('done', 0)}")
    print(f"    Failed:  {csv_counts.get('failed', 0)}")

    print(f"\n  Excel ({OUTPUT_XLSX}):")
    print(f"    Total rows: {len(excel_ids)}")

    # Find inconsistencies
    csv_done    = {r["q_id"] for r in rows
                   if r["status"] == "done"}
    csv_failed  = {r["q_id"] for r in rows
                   if r["status"] == "failed"}
    excel_set   = set(excel_ids.keys())

    in_excel_not_done  = excel_set - csv_done
    done_not_in_excel  = csv_done - excel_set
    failed_in_excel    = {q for q, s in excel_ids.items()
                          if "FAILED" in s.upper()}

    print(f"\n  Inconsistencies:")
    print(f"    In Excel but not done in CSV: "
          f"{len(in_excel_not_done)}")
    print(f"    Done in CSV but not in Excel: "
          f"{len(done_not_in_excel)}")
    print(f"    Failed rows in Excel:         "
          f"{len(failed_in_excel)}")

    if csv_failed:
        print(f"\n  Failed Q_IDs in CSV:")
        for r in rows:
            if r["status"] == "failed":
                print(f"    {r['q_id']:15} | "
                      f"{r['learning_objective'][:40]}")

    print("=" * 55)


def cmd_fix_failed():
    """Reset all failed rows in CSV and delete from Excel."""
    rows, fieldnames = load_csv()
    excel_ids        = load_excel_ids()

    # Find failed in CSV
    failed_ids = [r["q_id"] for r in rows
                  if r["status"] == "failed"]

    # Find failed in Excel
    excel_failed = [q for q, s in excel_ids.items()
                    if "FAILED" in s.upper()]

    all_to_reset = list(set(failed_ids + excel_failed))

    if not all_to_reset:
        print("[Maintain] No failed rows found.")
        return

    # Reset in CSV
    reset_count = 0
    for row in rows:
        if row["q_id"] in all_to_reset:
            row["status"]       = "pending"
            row["generated_at"] = ""
            reset_count += 1

    save_csv(rows, fieldnames)
    print(f"[CSV] Reset {reset_count} rows to pending")

    # Delete from Excel
    deleted = delete_from_excel(all_to_reset)
    print(f"[Excel] Deleted {deleted} failed rows")

    print(f"\n[Maintain] Fixed {len(all_to_reset)} rows:")
    for q_id in all_to_reset:
        print(f"  {q_id}")

    print(f"\nRun to regenerate:")
    print(f"  python dataset/generate_dataset.py --batch 15")


def cmd_reset(q_ids: list):
    """Reset specific Q_IDs to pending."""
    rows, fieldnames = load_csv()

    reset_count = 0
    for row in rows:
        if row["q_id"] in q_ids:
            row["status"]       = "pending"
            row["generated_at"] = ""
            reset_count += 1

    save_csv(rows, fieldnames)
    print(f"[CSV] Reset {reset_count} rows to pending")

    deleted = delete_from_excel(q_ids)
    print(f"[Excel] Deleted {deleted} rows")

    print(f"\n[Maintain] Reset complete: {q_ids}")


def cmd_reset_category(category: str):
    """Reset entire category to pending."""
    rows, fieldnames = load_csv()

    q_ids       = []
    reset_count = 0
    for row in rows:
        if row["category"] == category:
            q_ids.append(row["q_id"])
            row["status"]       = "pending"
            row["generated_at"] = ""
            reset_count += 1

    save_csv(rows, fieldnames)
    print(f"[CSV] Reset {reset_count} rows to pending "
          f"in category: {category}")

    deleted = delete_from_excel(q_ids)
    print(f"[Excel] Deleted {deleted} rows")


def cmd_sync():
    """Sync CSV with Excel — mark done in CSV if in Excel."""
    rows, fieldnames = load_csv()
    excel_ids        = load_excel_ids()

    sync_count = 0
    for row in rows:
        q_id   = row["q_id"]
        status = excel_ids.get(q_id, "")
        if q_id in excel_ids and \
                "FAILED" not in status.upper() and \
                row["status"] != "done":
            row["status"] = "done"
            sync_count += 1

    save_csv(rows, fieldnames)
    print(f"[Maintain] Synced {sync_count} rows "
          f"from Excel → CSV")


# ════════════════════════════════════════════════════════════
# Main
# ════════════════════════════════════════════════════════════

if __name__ == "__main__":
    parser = argparse.ArgumentParser(
        description="Dataset maintenance utility"
    )
    parser.add_argument(
        "--status", action="store_true",
        help="Show current dataset status"
    )
    parser.add_argument(
        "--fix-failed", action="store_true",
        help="Reset all failed rows in CSV and Excel"
    )
    parser.add_argument(
        "--reset", nargs="+", metavar="Q_ID",
        help="Reset specific Q_IDs e.g. --reset E-DE-010 E-DE-011"
    )
    parser.add_argument(
        "--reset-category", metavar="CATEGORY",
        help="Reset entire category e.g. easy_deep_learning"
    )
    parser.add_argument(
        "--sync", action="store_true",
        help="Sync CSV status with Excel"
    )

    args = parser.parse_args()

    if args.status:
        cmd_status()
    elif args.fix_failed:
        cmd_fix_failed()
    elif args.reset:
        cmd_reset(args.reset)
    elif args.reset_category:
        cmd_reset_category(args.reset_category)
    elif args.sync:
        cmd_sync()
    else:
        parser.print_help()
