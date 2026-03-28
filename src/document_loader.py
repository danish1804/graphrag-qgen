# ============================================================
# src/document_loader.py
# Loads PDF and PPTX files and returns clean text chunks
# ============================================================

import os
from pathlib import Path
from typing import List, Dict

from pypdf import PdfReader
from pptx import Presentation
from langchain.text_splitter import RecursiveCharacterTextSplitter


def load_pdf(file_path: str) -> str:
    """Extract raw text from a PDF file."""
    reader = PdfReader(file_path)
    text = ""
    for page in reader.pages:
        page_text = page.extract_text()
        if page_text:
            text += page_text + "\n"
    return text.strip()

def load_pptx(file_path: str) -> str:
    """Extract raw text from a PPTX file."""
    prs = Presentation(file_path)
    text = ""
    for slide_num, slide in enumerate(prs.slides, 1):
        slide_text = ""
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                slide_text += shape.text.strip() + "\n"
        if slide_text:
            text += f"[Slide {slide_num}]\n{slide_text}\n"
    return text.strip()



