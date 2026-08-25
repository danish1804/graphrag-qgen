# ============================================================
# scripts/corpus_to_excel.py
# Export the FULL question corpus to Excel, for your own review.
#
# This is NOT the rater workbook. It contains everything, including
# fields that are deliberately hidden from raters:
#   - the sample answer          (would give away solvability)
#   - the intended difficulty    (raters judge difficulty themselves)
#   - pair_id                    (reveals the matched-pair design)
#   - routing provenance and judge scores
#
# Never send this file to the raters. Use
# docs/human_evaluation_ROUND1.xlsx for that.
#
# Usage:
#   python scripts/corpus_to_excel.py
#   python scripts/corpus_to_excel.py --corpus outputs/question_corpus.json \
#       --out outputs/question_corpus.xlsx
# ============================================================
import os
import re
import sys
import json
import argparse

sys.path.append(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DEFAULT_IN = os.path.join(ROOT, "outputs", "question_corpus.json")
DEFAULT_OUT = os.path.join(ROOT, "outputs", "question_corpus.xlsx")

try:
    from scripts.inspect_corpus import dangling_refs
except Exception:                       # standalone fallback
    def dangling_refs(_):
        return []


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--corpus", default=DEFAULT_IN)
    ap.add_argument("--out", default=DEFAULT_OUT)
    args = ap.parse_args()

    if not os.path.isfile(args.corpus):
        sys.exit(f"[!] corpus not found: {args.corpus}")

    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    payload = json.load(open(args.corpus, encoding="utf-8"))
    qs = payload.get("questions", [])
    if not qs:
        sys.exit("[!] corpus contains no questions")

    ACCENT = "1F4E79"
    F = "Arial"
    thin = Side(style="thin", color="BFBFBF")
    BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

    wb = openpyxl.Workbook()

    # ── Summary ────────────────────────────────────────────
    ws = wb.active
    ws.title = "Summary"
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 32
    ws.column_dimensions["C"].width = 74

    t = ws.cell(row=1, column=2, value="Question Corpus — full export")
    t.font = Font(name=F, bold=True, size=15, color=ACCENT)
    w = ws.cell(row=2, column=2,
                value="INTERNAL — contains answers, intended difficulty and "
                      "pair ids. Do NOT send to raters.")
    w.font = Font(name=F, bold=True, size=11, color="C00000")

    flagged = [q for q in qs if dangling_refs(q.get("question"))]
    from collections import Counter
    rows = [
        ("Source file", os.path.basename(args.corpus)),
        ("Generated at", payload.get("generated_at", "")),
        ("Questions", len(qs)),
        ("Judge enabled at generation", str(payload.get("judge_enabled"))),
        ("Domain forced", str(payload.get("force_domain"))),
        ("Failures during generation", payload.get("failures", 0)),
        ("Not self-contained", f"{len(flagged)} ({len(flagged)/len(qs)*100:.0f}%)"),
        ("Matched pairs",
         f"{sum(1 for q in qs if q.get('pair_id'))} items"),
        ("", ""),
        ("By domain", ", ".join(f"{k} {v}" for k, v in
                                sorted(Counter(q.get('domain') for q in qs).items()))),
        ("By difficulty", ", ".join(f"{k} {v}" for k, v in
                                    sorted(Counter(q.get('difficulty') for q in qs).items()))),
        ("By question type", ", ".join(f"{k} {v}" for k, v in
                                       sorted(Counter(q.get('question_type') for q in qs).items()))),
    ]
    r = 4
    for label, val in rows:
        if label:
            a = ws.cell(row=r, column=2, value=label)
            a.font = Font(name=F, bold=True, size=11, color=ACCENT)
            b = ws.cell(row=r, column=3, value=val)
            b.font = Font(name=F, size=11)
            b.alignment = Alignment(wrap_text=True, vertical="top")
        r += 1

    # ── Questions ──────────────────────────────────────────
    qsh = wb.create_sheet("Questions")
    qsh.sheet_view.showGridLines = False
    cols = [
        ("question_id", 13), ("domain", 21), ("requested_domain", 21),
        ("difficulty", 12), ("question_type", 14), ("pair_id", 9),
        ("self_contained", 15), ("learning_outcome", 46),
        ("question", 84), ("sample_answer", 84),
        ("routing_confidence", 17), ("judge_question", 14),
    ]
    for i, (name, wd) in enumerate(cols, start=1):
        c = qsh.cell(row=1, column=i, value=name)
        c.font = Font(name=F, bold=True, size=11, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=ACCENT)
        c.alignment = Alignment(horizontal="center", vertical="center",
                                wrap_text=True)
        c.border = BORDER
        qsh.column_dimensions[get_column_letter(i)].width = wd
    qsh.row_dimensions[1].height = 28
    qsh.freeze_panes = "B2"

    ALT = PatternFill("solid", fgColor="F2F6FB")
    BAD = PatternFill("solid", fgColor="FCE4E4")

    def routing_summary(rc):
        """routing_confidence is a dict for auto-routed items and absent
        for pinned ones. Flatten to a readable cell."""
        if not rc:
            return "pinned"
        if not isinstance(rc, dict):
            return rc
        top = rc.get("top")
        margin = rc.get("margin")
        amb = rc.get("ambiguous")
        parts = []
        if top:
            parts.append(f"{top} {rc.get('top_score', '')}")
        if margin is not None:
            parts.append(f"margin {margin}")
        if amb:
            parts.append("AMBIGUOUS")
        return " | ".join(str(p) for p in parts) or "auto"

    for i, q in enumerate(qs):
        row = 2 + i
        dang = dangling_refs(q.get("question"))
        js = q.get("judge_scores") or {}
        vals = [
            q.get("question_id"), q.get("domain"), q.get("requested_domain"),
            q.get("difficulty"), q.get("question_type"), q.get("pair_id"),
            "NO — " + dang[0] if dang else "yes",
            q.get("learning_outcome"), q.get("question"),
            q.get("sample_answer"),
            routing_summary(q.get("routing_confidence")),
            js.get("question_overall"),
        ]
        for j, v in enumerate(vals, start=1):
            c = qsh.cell(row=row, column=j, value=v)
            c.font = Font(name=F, size=10)
            c.alignment = Alignment(vertical="top",
                                    wrap_text=j in (8, 9, 10))
            c.border = BORDER
            if dang:
                c.fill = BAD
            elif i % 2:
                c.fill = ALT

    os.makedirs(os.path.dirname(args.out) or ".", exist_ok=True)
    wb.save(args.out)
    print(f"Wrote {args.out}")
    print(f"  {len(qs)} questions, {len(flagged)} flagged as not self-contained")
    print("  INTERNAL FILE — do not send to raters.")


if __name__ == "__main__":
    main()
